#!/usr/bin/env python3
"""Render anonymous MEE DOCX with double spacing and continuous line numbering."""
from __future__ import annotations
import argparse
import re
from pathlib import Path
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

INLINE = re.compile(r"(\*\*[^*]+\*\*|(?<!\*)\*[^*]+\*(?!\*)|\x60[^\x60]+\x60)")
TABLE_SEPARATOR = re.compile(r"^\s*\|?\s*:?-{3,}:?\s*(\|\s*:?-{3,}:?\s*)+\|?\s*$")
URL_LINK = re.compile(r"\[([^\]]+)\]\((https?://[^)]+)\)")

def add_emphasis(paragraph, text: str) -> None:
    text = URL_LINK.sub(lambda m: f"{m.group(1)} ({m.group(2)})", text.strip())
    pos = 0
    for match in INLINE.finditer(text):
        if match.start() > pos:
            paragraph.add_run(text[pos:match.start()])
        token = match.group(0)
        run = paragraph.add_run(token[2:-2] if token.startswith("**") else token[1:-1])
        if token.startswith("**"):
            run.bold = True
        elif token.startswith("*"):
            run.italic = True
        else:
            run.font.name = "Courier New"
        pos = match.end()
    if pos < len(text):
        paragraph.add_run(text[pos:])

def format_base(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin, section.bottom_margin = Cm(2.5), Cm(2.5)
    section.left_margin, section.right_margin = Cm(3.2), Cm(2.5)
    section.header_distance, section.footer_distance = Cm(1), Cm(1)
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)
    normal.paragraph_format.line_spacing = 2
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.widow_control = True
    title = doc.styles["Title"]
    title.font.name = "Times New Roman"
    title.font.size = Pt(14)
    title.font.bold = True
    title.font.color.rgb = RGBColor(0, 0, 0)
    title_ppr = title._element.get_or_add_pPr()
    for border in title_ppr.findall(qn("w:pBdr")):
        title_ppr.remove(border)
    title.paragraph_format.line_spacing = 2
    title.paragraph_format.keep_with_next = True
    title.paragraph_format.space_after = Pt(12)
    for level in range(1, 4):
        style = doc.styles[f"Heading {level}"]
        style.font.name = "Times New Roman"
        style.font.size = Pt(12)
        style.font.bold = True
        style.font.italic = False
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.line_spacing = 2
        style.paragraph_format.space_before = Pt(10 if level <= 2 else 6)
        style.paragraph_format.space_after = Pt(0)
        style.paragraph_format.keep_with_next = True
    for name in ("List Bullet", "List Number"):
        style = doc.styles[name]
        style.font.name = "Times New Roman"
        style.font.size = Pt(12)
        style.paragraph_format.line_spacing = 2
        style.paragraph_format.space_after = Pt(0)
    # Native Word continuous line numbers, not manually simulated labels.
    sectpr = section._sectPr
    for old in sectpr.findall(qn("w:lnNumType")):
        sectpr.remove(old)
    line = OxmlElement("w:lnNumType")
    for key, value in (("countBy", "1"), ("start", "1"), ("restart", "continuous"), ("distance", "360")):
        line.set(qn("w:" + key), value)
    sectpr.append(line)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.style = "Normal"
    footer.paragraph_format.line_spacing = 1
    footer.add_run("Page ")
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    footer._p.append(fld)
    doc.core_properties.author = ""
    doc.core_properties.last_modified_by = ""
    doc.core_properties.keywords = ""
    doc.core_properties.comments = ""

def add_table(doc: Document, rows: list[str]) -> None:
    split = lambda row: [x.strip() for x in row.strip().strip("|").split("|")]
    parsed = [split(row) for row in rows if not TABLE_SEPARATOR.match(row)]
    if not parsed:
        return
    table = doc.add_table(rows=len(parsed), cols=max(map(len, parsed)))
    table.style = "Table Grid"
    for i, values in enumerate(parsed):
        for j, value in enumerate(values):
            p = table.cell(i, j).paragraphs[0]
            p.paragraph_format.line_spacing = 2
            add_emphasis(p, value)
            if i == 0:
                for run in p.runs:
                    run.bold = True

def render_manuscript(source: Path, output: Path) -> None:
    text = source.read_text(encoding="utf-8")
    doc = Document()
    format_base(doc)
    lines = text.splitlines()
    i = 0
    while i < len(lines):
        line = lines[i].strip()
        if not line or line == "---":
            i += 1
            continue
        if line.startswith("|") and i + 1 < len(lines) and TABLE_SEPARATOR.match(lines[i + 1]):
            rows = [line]
            i += 1
            while i < len(lines) and lines[i].lstrip().startswith("|"):
                rows.append(lines[i])
                i += 1
            add_table(doc, rows)
            continue
        if line.startswith("# "):
            p = doc.add_paragraph(style="Title")
            add_emphasis(p, line[2:])
        elif line.startswith("### "):
            p = doc.add_paragraph(style="Heading 2")
            add_emphasis(p, line[4:])
        elif line.startswith("## "):
            p = doc.add_paragraph(style="Heading 1")
            add_emphasis(p, line[3:])
        elif re.match(r"^\s*[-*]\s+", line):
            p = doc.add_paragraph(style="List Bullet")
            add_emphasis(p, re.sub(r"^\s*[-*]\s+", "", line))
        else:
            p = doc.add_paragraph(style="Normal")
            if line.startswith(("Aarts,", "Dormann,", "Fisher,", "Galipaud,", "Getz,", "Roberts,", "Strobl,", "Zbinden,")):
                p.paragraph_format.left_indent = Cm(0.5)
                p.paragraph_format.first_line_indent = Cm(-0.5)
            add_emphasis(p, line)
        i += 1
    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)

def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--source", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    render_manuscript(args.source, args.output)
    print(args.output)

if __name__ == "__main__":
    main()
