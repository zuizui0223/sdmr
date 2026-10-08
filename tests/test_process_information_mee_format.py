"""Check that generated MEE main-manuscript DOCX is anonymous and legible by structure."""
from __future__ import annotations
import subprocess
import sys
import zipfile
from pathlib import Path
from xml.etree import ElementTree

from docx import Document

ROOT=Path(__file__).resolve().parents[1]
SOURCE=ROOT/"manuscript"/"PROCESS_INFORMATION_MEE_SUBMISSION_V1.md"
SCRIPT=ROOT/"scripts"/"build_process_information_mee_manuscript_docx.py"
NS={"w":"http://schemas.openxmlformats.org/wordprocessingml/2006/main"}

def test_mee_submission_docx_has_numbering_spacing_and_no_author(tmp_path):
    output=tmp_path/"review_anonymous.docx"
    subprocess.run([sys.executable,str(SCRIPT),"--source",str(SOURCE),
                    "--output",str(output)],check=True,capture_output=True,text=True)
    assert output.stat().st_size>10000
    doc=Document(output)
    text="\n".join(p.text for p in doc.paragraphs)
    source=SOURCE.read_text(encoding="utf-8")
    for phrase in ("71 of 80", "0/700", "16/300", "10/50",
                   "Strobl et al. 2008", "Fisher et al. 2019",
                   "No head-to-head benchmark", "References"):
        assert phrase in source
        assert phrase in text
    assert doc.core_properties.author == ""
    assert doc.core_properties.last_modified_by == ""
    assert "zuizui0223" not in text.lower()
    assert "github.com/zuizui0223" not in text.lower()
    assert doc.sections[0].page_width is not None
    normal=doc.styles["Normal"]
    assert normal.paragraph_format.line_spacing == 2
    assert "Times New Roman" == normal.font.name
    assert doc.styles["Title"].font.color.rgb is not None
    assert str(doc.styles["Title"].font.color.rgb) == "000000"
    assert doc.styles["Title"]._element.pPr.find("{"+NS["w"]+"}pBdr") is None
    assert str(doc.styles["Heading 1"].font.color.rgb) == "000000"
    assert len(doc.paragraphs)>60

    with zipfile.ZipFile(output) as z:
        main=ElementTree.fromstring(z.read("word/document.xml"))
        sect=main.find(".//w:sectPr",NS)
        assert sect is not None
        lines=sect.find("w:lnNumType",NS)
        assert lines is not None
        assert lines.attrib.get("{"+NS["w"]+"}restart")=="continuous"
        assert lines.attrib.get("{"+NS["w"]+"}countBy")=="1"
        footer_names=[x for x in z.namelist() if x.startswith("word/footer") and x.endswith(".xml")]
        assert footer_names
        assert any(b"PAGE" in z.read(f) for f in footer_names)
        core=z.read("docProps/core.xml").decode("utf-8")
        assert "ZHANG" not in core and "zuizui0223" not in core

def test_mee_docx_builder_handles_bold_italic_bullets_and_table(tmp_path):
    source=tmp_path/"fixture.md"
    source.write_text(
        "# Scientific title\n\n## Abstract\n\n1. **Finding** from *worlds*.\n\n"
        "- a bullet with **emphasis**\n\n"
        "| A | B |\n|---|---|\n| pass | fail |\n",
        encoding="utf-8",
    )
    output=tmp_path/"fixture.docx"
    subprocess.run([sys.executable,str(SCRIPT),"--source",str(source),
                    "--output",str(output)],check=True)
    doc=Document(output)
    assert doc.tables and doc.tables[0].cell(1,0).text=="pass"
    assert any("Finding" in p.text for p in doc.paragraphs)
    assert any(r.bold for p in doc.paragraphs for r in p.runs if "Finding" in r.text)
